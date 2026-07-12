"""Autofit: normAutofit with PowerPoint-style precomputed fontScale /
lnSpcReduction (stored values are deterministic to honor), plus a no-autofit
overflow control box."""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

from _gen_common import blank_slide, body_autofit, grid_boxes, new_deck, save
from pptx.dml.color import RGBColor
from pptx.util import Pt

FONT = "Calibri"
LONG = ("Shrink-to-fit text keeps reducing the font scale until the whole "
        "paragraph fits inside its box. ") * 4


def filled_box(slide, cell, text, size=20):
    tb = slide.shapes.add_textbox(*cell)
    tf = tb.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = FONT
    tb.line.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
    tb.line.width = Pt(0.75)
    return tf


def main():
    prs = new_deck()
    slide = blank_slide(prs)
    cells = list(grid_boxes(4, cols=2))

    body_autofit(filled_box(slide, cells[0], LONG), font_scale=62500)
    body_autofit(filled_box(slide, cells[1], LONG),
                 font_scale=40000, lnspc_reduction=20000)
    body_autofit(filled_box(slide, cells[2], "Fits fine, scale 100%.", 16),
                 font_scale=100000)
    filled_box(slide, cells[3], LONG)  # no autofit: overflow control

    save(prs)


if __name__ == "__main__":
    main()
